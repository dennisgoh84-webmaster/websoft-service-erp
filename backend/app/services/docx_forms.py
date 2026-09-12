"""Word (.docx) versions of the printable forms -- same content as the
matching print page (e.g. frontend/src/pages/InvoicePrintPage.tsx for
invoice_to_docx), built with python-docx so it opens as a normal,
editable Word document. Each form gets its own small function here
rather than a generic templating layer, since a one-record form's
layout is specific to what it's showing.
"""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.models.billing import Invoice
from app.models.core import Company
from app.models.company_individuals import CompanyIndividual
from app.models.job_orders import JobOrder
from app.models.payables import PurchaseOrder, SupplierPayment
from app.models.payments import Payment
from app.models.quotations import Quotation
from app.models.service_records import ServiceRecord


def invoice_to_docx(invoice: Invoice, customer: CompanyIndividual, company: Company) -> bytes:
    doc = Document()

    header = doc.add_paragraph()
    header.add_run(company.name).bold = True
    if company.address:
        doc.add_paragraph(company.address)
    if company.phone:
        doc.add_paragraph(f"Tel: {company.phone}")
    if company.website:
        doc.add_paragraph(company.website)
    if company.uen:
        doc.add_paragraph(f"Business Reg# {company.uen}")
    if company.gst_registration_no:
        doc.add_paragraph(f"GST Reg# {company.gst_registration_no}")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TAX INVOICE")
    run.bold = True
    run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.add_run(f"{invoice.invoice_number}\n").bold = True
    meta.add_run(f"Issued: {invoice.issued_at.date().isoformat()}\n")
    if invoice.due_date:
        meta.add_run(f"Due: {invoice.due_date.isoformat()}\n")

    doc.add_paragraph().add_run("Bill To").italic = True
    bill_to = doc.add_paragraph()
    bill_to.add_run(customer.name + "\n").bold = True
    if customer.uen:
        bill_to.add_run(f"UEN: {customer.uen}\n")
    if customer.contact_person:
        bill_to.add_run(f"Contact Person: {customer.contact_person}\n")
    if customer.billing_email:
        bill_to.add_run(f"Contact Email: {customer.billing_email}\n")
    if customer.phone:
        bill_to.add_run(f"Contact No: {customer.phone}\n")
    address = ", ".join(
        filter(None, [customer.address_line1, customer.address_line2, customer.address_city, customer.address_country])
    )
    if address:
        bill_to.add_run(address)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Description"
    hdr[1].text = "Qty"
    hdr[2].text = "Unit Price ($)"
    hdr[3].text = "Amount ($)"
    row = table.add_row().cells
    row[0].text = invoice.description
    row[1].text = "1.00"
    row[2].text = f"{invoice.amount_sgd:.2f}"
    row[3].text = f"{invoice.amount_sgd:.2f}"

    doc.add_paragraph()
    totals = doc.add_table(rows=3, cols=2)
    for i, (label, value) in enumerate(
        [
            ("Subtotal", f"{invoice.amount_sgd:.2f}"),
            (f"Tax {invoice.gst_rate}% ({invoice.tax_code})", f"{invoice.gst_amount_sgd:.2f}"),
            ("Grand Total (SGD)", f"{invoice.total_amount_sgd:.2f}"),
        ]
    ):
        cells = totals.rows[i].cells
        cells[0].text = label
        cells[1].text = value
        if label.startswith("Grand Total"):
            for cell in cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def quotation_to_docx(quotation: Quotation, customer: CompanyIndividual, company: Company) -> bytes:
    doc = Document()

    header = doc.add_paragraph()
    header.add_run(company.name).bold = True
    if company.address:
        doc.add_paragraph(company.address)
    if company.phone:
        doc.add_paragraph(f"Tel: {company.phone}")
    if company.website:
        doc.add_paragraph(company.website)
    if company.uen:
        doc.add_paragraph(f"Business Reg# {company.uen}")
    if company.gst_registration_no:
        doc.add_paragraph(f"GST Reg# {company.gst_registration_no}")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("QUOTATION")
    run.bold = True
    run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.add_run(f"{quotation.quotation_number}\n").bold = True
    meta.add_run(f"Date: {quotation.quotation_date.isoformat()}\n")
    if quotation.valid_until:
        meta.add_run(f"Valid Until: {quotation.valid_until.isoformat()}\n")

    doc.add_paragraph().add_run("To").italic = True
    bill_to = doc.add_paragraph()
    bill_to.add_run(customer.name + "\n").bold = True
    if customer.uen:
        bill_to.add_run(f"UEN: {customer.uen}\n")
    if customer.contact_person:
        bill_to.add_run(f"Contact Person: {customer.contact_person}\n")
    if customer.billing_email:
        bill_to.add_run(f"Contact Email: {customer.billing_email}\n")
    if customer.phone:
        bill_to.add_run(f"Contact No: {customer.phone}\n")
    address = ", ".join(
        filter(None, [customer.address_line1, customer.address_line2, customer.address_city, customer.address_country])
    )
    if address:
        bill_to.add_run(address)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Description"
    hdr[1].text = "UoM"
    hdr[2].text = "Qty"
    hdr[3].text = "Unit Price ($)"
    hdr[4].text = "Amount ($)"
    for line in quotation.lines:
        row = table.add_row().cells
        row[0].text = line.description
        row[1].text = line.unit_of_measure or ""
        row[2].text = f"{float(line.quantity):.2f}"
        row[3].text = f"{float(line.unit_price_sgd):.2f}"
        row[4].text = f"{float(line.line_total_sgd):.2f}"

    doc.add_paragraph()
    totals = doc.add_table(rows=3, cols=2)
    for i, (label, value) in enumerate(
        [
            ("Subtotal", f"{float(quotation.amount_sgd):.2f}"),
            (f"Tax {float(quotation.gst_rate)}% ({quotation.tax_code})", f"{float(quotation.gst_amount_sgd):.2f}"),
            ("Grand Total (SGD)", f"{float(quotation.total_amount_sgd):.2f}"),
        ]
    ):
        cells = totals.rows[i].cells
        cells[0].text = label
        cells[1].text = value
        if label.startswith("Grand Total"):
            for cell in cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True

    if quotation.notes:
        doc.add_paragraph()
        notes = doc.add_paragraph()
        notes.add_run("Notes: ").bold = True
        notes.add_run(quotation.notes)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def receipt_to_docx(
    payment: Payment, customer: CompanyIndividual, company: Company, invoice_numbers: dict
) -> bytes:
    doc = Document()

    header = doc.add_paragraph()
    header.add_run(company.name).bold = True
    if company.address:
        doc.add_paragraph(company.address)
    if company.phone:
        doc.add_paragraph(f"Tel: {company.phone}")
    if company.uen:
        doc.add_paragraph(f"Business Reg# {company.uen}")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("OFFICIAL RECEIPT")
    run.bold = True
    run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.add_run(f"{payment.voucher_number}\n").bold = True
    meta.add_run(f"Date: {payment.payment_date.isoformat()}\n")
    meta.add_run(f"Method: {payment.method.value}\n")
    if payment.reference:
        meta.add_run(f"Reference: {payment.reference}\n")

    doc.add_paragraph().add_run("Received From").italic = True
    from_p = doc.add_paragraph()
    from_p.add_run(customer.name + "\n").bold = True
    if customer.uen:
        from_p.add_run(f"UEN: {customer.uen}")

    doc.add_paragraph()
    total = doc.add_paragraph()
    total.add_run(f"Amount Received: SGD {float(payment.amount_sgd):.2f}").bold = True

    if payment.allocations:
        doc.add_paragraph().add_run("Applied To").italic = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Invoice"
        hdr[1].text = "Amount ($)"
        for a in payment.allocations:
            row = table.add_row().cells
            row[0].text = invoice_numbers.get(a.invoice_id, "")
            row[1].text = f"{float(a.amount_sgd):.2f}"

    unallocated = float(payment.unallocated_sgd)
    if unallocated > 0:
        doc.add_paragraph()
        doc.add_paragraph(f"Unallocated (on account): SGD {unallocated:.2f}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def purchase_order_to_docx(po: PurchaseOrder, supplier: CompanyIndividual, company: Company) -> bytes:
    """Same layout as frontend/src/pages/PurchaseOrderPrintPage.tsx --
    also what "Email PO" (2026-09-12) converts to PDF and attaches."""
    doc = Document()

    header = doc.add_paragraph()
    header.add_run(company.name).bold = True
    if company.address:
        doc.add_paragraph(company.address)
    if company.phone:
        doc.add_paragraph(f"Tel: {company.phone}")
    if company.uen:
        doc.add_paragraph(f"Business Reg# {company.uen}")
    if company.gst_registration_no:
        doc.add_paragraph(f"GST Reg# {company.gst_registration_no}")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PURCHASE ORDER")
    run.bold = True
    run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.add_run(f"{po.po_number}\n").bold = True
    meta.add_run(f"Date: {po.order_date.isoformat()}\n")
    meta.add_run(f"Status: {po.status.value.replace('_', ' ').title()}\n")

    doc.add_paragraph().add_run("Supplier").italic = True
    to_p = doc.add_paragraph()
    to_p.add_run(supplier.name + "\n").bold = True
    if supplier.gst_registration_no:
        to_p.add_run(f"GST Reg# {supplier.gst_registration_no}\n")
    if supplier.billing_email:
        to_p.add_run(f"Email: {supplier.billing_email}\n")
    if supplier.phone:
        to_p.add_run(f"Tel: {supplier.phone}\n")
    address = ", ".join(
        filter(None, [supplier.address_line1, supplier.address_line2, supplier.address_city, supplier.address_country])
    )
    if address:
        to_p.add_run(address)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Description"
    hdr[1].text = "Amount ($)"
    row = table.add_row().cells
    row[0].text = po.description
    row[1].text = f"{float(po.amount_sgd):.2f}"

    doc.add_paragraph()
    totals = doc.add_table(rows=2, cols=2)
    for i, (label, value) in enumerate(
        [
            ("GST", f"{float(po.gst_amount_sgd):.2f}"),
            ("Grand Total (SGD)", f"{float(po.total_amount_sgd):.2f}"),
        ]
    ):
        cells = totals.rows[i].cells
        cells[0].text = label
        cells[1].text = value
        if label.startswith("Grand Total"):
            for cell in cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True

    doc.add_paragraph()
    doc.add_paragraph("Please confirm receipt of this purchase order and quote the PO number "
                       "above on your invoice.")
    doc.add_paragraph()
    doc.add_paragraph("Authorised by: ______________________________")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def payment_voucher_to_docx(
    payment: SupplierPayment, supplier: CompanyIndividual, company: Company, bill_numbers: dict
) -> bytes:
    doc = Document()

    header = doc.add_paragraph()
    header.add_run(company.name).bold = True
    if company.address:
        doc.add_paragraph(company.address)
    if company.phone:
        doc.add_paragraph(f"Tel: {company.phone}")
    if company.uen:
        doc.add_paragraph(f"Business Reg# {company.uen}")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PAYMENT VOUCHER")
    run.bold = True
    run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.add_run(f"{payment.voucher_number}\n").bold = True
    meta.add_run(f"Date: {payment.payment_date.isoformat()}\n")
    meta.add_run(f"Method: {payment.method}\n")
    if payment.reference:
        meta.add_run(f"Reference: {payment.reference}\n")

    doc.add_paragraph().add_run("Paid To").italic = True
    to_p = doc.add_paragraph()
    to_p.add_run(supplier.name + "\n").bold = True
    if supplier.gst_registration_no:
        to_p.add_run(f"GST Reg# {supplier.gst_registration_no}")

    doc.add_paragraph()
    total = doc.add_paragraph()
    total.add_run(f"Amount Paid: SGD {float(payment.amount_sgd):.2f}").bold = True

    if payment.allocations:
        doc.add_paragraph().add_run("Applied To").italic = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Bill"
        hdr[1].text = "Amount ($)"
        for a in payment.allocations:
            row = table.add_row().cells
            row[0].text = bill_numbers.get(a.supplier_invoice_id, "")
            row[1].text = f"{float(a.amount_sgd):.2f}"

    unallocated = float(payment.unallocated_sgd)
    if unallocated > 0:
        doc.add_paragraph()
        doc.add_paragraph(f"Unallocated: SGD {unallocated:.2f}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def service_record_to_docx(record: ServiceRecord, job_order: JobOrder, customer: CompanyIndividual, company: Company) -> bytes:
    """Same layout as frontend/src/pages/ServiceRecordPrintPage.tsx --
    also what "Email" (2026-09-12) converts to PDF and attaches."""
    doc = Document()

    header = doc.add_paragraph()
    header.add_run(company.name).bold = True
    if company.address:
        doc.add_paragraph(company.address)
    if company.phone:
        doc.add_paragraph(f"Tel: {company.phone}")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SERVICE RECORD")
    run.bold = True
    run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.add_run(f"{record.service_record_number}\n").bold = True
    meta.add_run(f"Job Order: {job_order.job_order_number} -- {job_order.subject}\n")
    meta.add_run(f"Work Date: {record.work_date.isoformat()}\n")
    meta.add_run(f"Status: {record.status.value.title()}\n")

    doc.add_paragraph().add_run("Company / Individual").italic = True
    to_p = doc.add_paragraph()
    to_p.add_run(customer.name)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Field"
    hdr[1].text = "Value"
    rows_data = [
        ("Time logged (raw)", f"{record.raw_minutes} min"),
        ("Time logged (rounded, SRV-007)", f"{record.rounded_minutes} min"),
        ("Completion", "Completed" if record.completion_status.value == "C" else "Not yet completed -- another visit expected"),
        ("After hours / weekend / holiday", "Yes" if record.is_after_hours else "No"),
    ]
    if record.deducted_minutes is not None:
        rows_data.append(("Minutes deducted from contract", f"{record.deducted_minutes} min"))
    for label, value in rows_data:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value

    doc.add_paragraph()
    doc.add_paragraph("Signature & Company Stamp: ______________________________")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def statement_to_docx(statement, customer: CompanyIndividual, company: Company) -> bytes:
    """`statement` is a CompanyIndividualStatement (see app/schemas/schemas.py) --
    accepted duck-typed rather than imported, so this services module
    doesn't take a dependency on the API schema layer. Same layout as
    frontend/src/pages/StatementPrintPage.tsx; also what "Email" converts
    to PDF and attaches (2026-09-12)."""
    doc = Document()

    header = doc.add_paragraph()
    header.add_run(company.name).bold = True
    if company.address:
        doc.add_paragraph(company.address)
    if company.gst_registration_no:
        doc.add_paragraph(f"GST Reg# {company.gst_registration_no}")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("STATEMENT OF ACCOUNTS")
    run.bold = True
    run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.add_run(f"{customer.name}\n").bold = True
    meta.add_run(f"As at: {statement.as_at.isoformat()}\n")
    if statement.payment_terms_days is not None:
        meta.add_run(f"Payment terms: Net {statement.payment_terms_days} days\n")

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Invoice"
    hdr[1].text = "Issued"
    hdr[2].text = "Due"
    hdr[3].text = "Total ($)"
    hdr[4].text = "Outstanding ($)"
    for line in statement.lines:
        row = table.add_row().cells
        row[0].text = line.invoice_number + (" (disputed)" if line.is_disputed else "")
        row[1].text = line.issued_on.isoformat()
        row[2].text = line.due_date.isoformat() if line.due_date else "-"
        row[3].text = f"{line.total_amount_sgd:.2f}"
        row[4].text = f"{line.outstanding_sgd:.2f}"

    doc.add_paragraph()
    totals = doc.add_paragraph()
    totals.add_run(f"Total Outstanding: SGD {statement.total_outstanding_sgd:.2f}").bold = True
    if statement.unallocated_credit_sgd > 0:
        doc.add_paragraph(f"Unallocated credit on account: SGD {statement.unallocated_credit_sgd:.2f}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
